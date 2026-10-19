import json
import os
from pathlib import Path
import subprocess
import sys

from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def test_docx_rules_run_without_web_or_model_sdks(tmp_path):
    paper = tmp_path / "paper.docx"
    doc = Document()
    doc.add_paragraph("A result with one source [1].")
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Smith A. Stellar spectra[J]. Astronomy, 2020, 10: 1-5.")
    doc.save(paper)
    evidence = tmp_path / "evidence.json"
    report = tmp_path / "report.md"
    runner = '''
import builtins, runpy, socket, sys
original_import = builtins.__import__
def local_import(name, *args, **kwargs):
    if name.split('.')[0] in {'fastapi', 'uvicorn', 'dashscope', 'langchain', 'langchain_openai', 'langchain_community', 'semanticscholar', 'crossref'}:
        raise ImportError('Not installed in this local-only runtime: ' + name)
    return original_import(name, *args, **kwargs)
builtins.__import__ = local_import
def no_network(*args, **kwargs):
    raise AssertionError('Document checking must stay local')
socket.create_connection = no_network
socket.socket.connect = no_network
sys.argv = sys.argv[1:]
runpy.run_path(sys.argv[0], run_name='__main__')
'''
    env = os.environ.copy()
    env.pop("MINERU_API_KEY", None)
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        [sys.executable, "-B", "-c", runner,
         str(ROOT / "scripts/progressive_papercheck.py"), str(paper),
         "--mode", "full", "--out", str(evidence), "--report", str(report)],
        env=env, capture_output=True, text=True, encoding="utf-8",
    )
    assert completed.returncode == 0, completed.stderr
    events = [json.loads(line) for line in completed.stdout.splitlines()]
    assert any(event["event"] == "rules_check_complete" for event in events)
    assert events[-1]["event"] == "papercheck_complete"
    payload = json.loads(evidence.read_text(encoding="utf-8"))
    assert payload["citation_count"] == 1
    assert payload["missing_citations"] == []
    assert "## Rules Summary" in report.read_text(encoding="utf-8")
